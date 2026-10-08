"""Tests fuer die Kette  PEP-PDF -> Standard-Excel -> excel_to_aasx.py -> AASX.

Aufruf (im Projektordner):  python -m pytest tests -q

Geprueft wird:
  - epd/<serie>/pep_<serie>.py lesen die PDFs und erzeugen exakt die Standard-Excel,
    die im Projekt liegt (Regression); Tabellen werden ohne Seitenvorgabe gefunden
  - Ankerwerte, von Hand aus den PDFs abgelesen, landen unveraendert in der AASX
  - die Plausibilitaetspruefung (Total = Summe der Module) erkennt verrutschte Spalten
  - leere Vorlage wird abgelehnt, befuellte Excel (221, 2273) angenommen
  - je Produkt eine AASX, aas-core3 ohne Befunde ausser AASd-120
  - ALLE Indikatorwerte der AASX = Blatt Ergebnisse (Referenz) bzw.
    Ergebnisse x Faktoren (Varianten, 3 signifikante Stellen), unabhaengig nachgerechnet
  - Vokabular des Templates (Assessed/NotAssessed, Declared/NotDeclared,
    Class, ThirdPartyVerification), kein Wert bei NotDeclared
  - Fehlerfaelle in der Excel fuehren zu Exit-Code 2 und keiner AASX
"""

import importlib.util
import shutil
import zipfile
from pathlib import Path

import pandas as pd
import pytest
from lxml import etree
from openpyxl import load_workbook

import epd_excel            # skripte/ liegt ueber conftest.py im Suchpfad
import excel_to_aasx
import pep_bausteine
from pdf_grid import open_pdf

PROJEKT = Path(__file__).resolve().parents[1]
EPD = PROJEKT / "epd"
TEMPLATE = PROJEKT / "vorlagen" / "epd-type-iii-submodel-template.aasx"
NS = "{https://admin-shell.io/aas/3/0}"


def pep_konfig(serie):
    """KONFIG aus epd/<serie>/pep_<serie>.py laden."""
    spec = importlib.util.spec_from_file_location(f"pep_{serie}", EPD / serie / f"pep_{serie}.py")
    modul = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modul)
    return modul.KONFIG


FAELLE = {"221": ("EPD_221.xlsx", pep_konfig("221"), 9), "2273": ("EPD_2273.xlsx", pep_konfig("2273"), 5)}

# Von Hand aus den PDFs abgelesen (S. 9, Zeile "Climate change - total"; Varianten:
# Referenzwert x Faktor von S. 14), Spalten A1-3, A4, A5, B1-7, C1-4, D
ANKER_GWP = {
    "221-422": [1.24e-2, 5.70e-4, 6.80e-4, 7.02e-2, 1.69e-3, -5.80e-4],
    "2273-202": [5.04e-3, 2.90e-4, 5.70e-4, 6.13e-2, 7.00e-4, -3.70e-4],
}
STUFEN = ["A1-A3", "A4", "A5", "B6", "C2-C4", "D"]


def lname(e):
    return etree.QName(e).localname


def lies_aasx(pfad):
    with zipfile.ZipFile(pfad) as z:
        namen = z.namelist()
        baum = etree.fromstring(z.read("aasx/data.xml"))
    sm = baum.find(f"{NS}submodels/{NS}submodel")
    return baum, sm, namen


def kinder(el):
    c = el.find(NS + "value") if lname(el) != "submodel" else el.find(NS + "submodelElements")
    return {x.findtext(NS + "idShort"): x for x in (c if c is not None else [])}


def werte_aus_aasx(sm):
    out, status = {}, []
    for e in kinder(kinder(sm)["EnvironmentalResults"]).values():
        k = kinder(e)
        code = k["indicatorCode"].findtext(NS + "value")
        for sv in k["stageValues"].find(NS + "value"):
            s = {x.findtext(NS + "idShort"): x.findtext(NS + "value") for x in sv.find(NS + "value")}
            status.append(s)
            if s.get("value") is not None:
                out[(code, s["stageCode"])] = float(s["value"])
    return out, status


def blatt(xlsx, name):
    """Listenblatt mit pandas lesen (Zeile 2 = Hinweiszeile), unabhaengig von epd_excel."""
    df = pd.read_excel(xlsx, name, skiprows=[1], dtype=str)
    df.columns = [c.replace(" *", "") for c in df.columns]
    return df.dropna(how="all")


def erwartete_werte(xlsx, artikel):
    """AASX-Werte unabhaengig aus der Standard-Excel nachrechnen (ohne Total)."""
    prod = blatt(xlsx, "Produkte")
    ref = prod.loc[prod["Referenzprodukt"] == "ja", "Artikelnummer"].iloc[0]
    erg = blatt(xlsx, "Ergebnisse")
    erg = erg[(erg["Artikelnummer"] == ref) & (erg["stageCode"] != "Total")]
    basis = {(r["indicatorCode"], r["stageCode"]): float(r["Wert"]) for _, r in erg.iterrows()}
    if artikel == ref:
        return basis
    fak = blatt(xlsx, "Faktoren")
    fak = fak[fak["Artikelnummer"] == artikel]
    f = {(r["indicatorCode"], r["stageCode"]): float(r["Faktor"]) for _, r in fak.iterrows()}
    return {k: float(f"{w * f[k]:.3g}") for k, w in basis.items()}


def inhalt(xlsx):
    roh = epd_excel.lese_arbeitsmappe(xlsx)
    return ({k: v for k, (v, _) in roh["epd"].items()},
            {b: [{k: v for k, v in z.items() if not k.startswith("_")} for z in zeilen]
             for b, zeilen in roh["listen"].items()})


@pytest.fixture(scope="module")
def arbeitsordner(tmp_path_factory):
    """Je EPD das PDF (und Produktbilder) in einen Arbeitsordner kopieren und das PEP-Skript
    dort laufen lassen - das Projekt selbst bleibt unveraendert."""
    d = tmp_path_factory.mktemp("epd")
    for serie, (_, konfig, _) in FAELLE.items():
        (d / serie).mkdir()
        shutil.copy(EPD / serie / konfig["pdf"], d / serie / konfig["pdf"])
        if (EPD / serie / "produktbilder").is_dir():
            shutil.copytree(EPD / serie / "produktbilder", d / serie / "produktbilder")
        assert pep_bausteine.pep_zu_vorlage(konfig, d / serie), f"PEP {serie}: Standard-Excel nicht erzeugt"
    epd_excel.schreibe_arbeitsmappe(d / "leer.xlsx")
    return d


def excel(ordner, serie):
    return ordner / serie / FAELLE[serie][0]


@pytest.fixture(scope="module")
def erzeugt(arbeitsordner):
    for serie in FAELLE:
        rc = excel_to_aasx.main([str(excel(arbeitsordner, serie)), "--out", str(arbeitsordner / "out" / serie),
                                 "--template", str(TEMPLATE)])
        assert rc == 0, f"excel_to_aasx {serie}: Exit-Code {rc}"
    return arbeitsordner / "out"


@pytest.mark.parametrize("serie", FAELLE)
def test_pdf_ergibt_die_excel_im_projekt(arbeitsordner, serie):
    """Regression: PDF -> Standard-Excel liefert genau die Datei, die im Projekt liegt."""
    assert inhalt(excel(arbeitsordner, serie)) == inhalt(excel(EPD, serie))


@pytest.mark.parametrize("serie, faktor_seiten", [("221", [13, 14]), ("2273", [13])])
def test_tabellen_werden_ohne_seitenvorgabe_gefunden(serie, faktor_seiten):
    s = pep_bausteine.finde_seiten(open_pdf(EPD / serie / FAELLE[serie][1]["pdf"]))
    assert (s["produkte"], s["setup"], s["indikatoren"], s["faktoren"]) == (3, 7, [8, 9, 10], faktor_seiten)


def test_summenpruefung_erkennt_verrutschte_spalte():
    werte = {"A1-3": 1.0, "A4": 0.1, "A5": 0.2, "B1-7": 5.0, "C1-4": 0.3, "Total": 6.6, "D": -0.1}
    zeilen = [{"indikator": "Climate change - total", "einheit": "kg", "spalte": s, "wert": w}
              for s, w in werte.items()]
    katalog = {"Climate change - total": "GWP-total"}
    prod = [{"artikel": "1-1", "serie": "x", "masse_g": 1.0}]
    fehler, _ = pep_bausteine.pruefe_extraktion(zeilen, [], prod, "1-1", katalog, "", "1")
    assert fehler == []
    zeilen[3]["wert"], zeilen[4]["wert"] = 0.3, 0.3            # B1-7 mit falschem Wert
    fehler, _ = pep_bausteine.pruefe_extraktion(zeilen, [], prod, "1-1", katalog, "", "1")
    assert any("Summe der Module" in f for f in fehler)


def test_ankerwerte_aus_dem_pdf(erzeugt):
    for artikel, werte in ANKER_GWP.items():
        serie = artikel.split("-")[0]
        _, sm, _ = lies_aasx(erzeugt / serie / f"wago-{artikel}-epd-submodel.aasx")
        aasx, _ = werte_aus_aasx(sm)
        assert [aasx[("GWP-total", s)] for s in STUFEN] == werte


def test_leere_vorlage_wird_abgelehnt(arbeitsordner):
    _, fehler, _ = epd_excel.pruefe(epd_excel.lese_arbeitsmappe(arbeitsordner / "leer.xlsx"))
    assert any("Registriernummer" in f for f in fehler)
    assert any("kein Produkt" in f for f in fehler)
    assert any("Module!" in f and "stageStatus" in f for f in fehler)


def test_leere_vorlage_hat_katalog_und_listen(arbeitsordner):
    wb = load_workbook(arbeitsordner / "leer.xlsx")
    assert wb.sheetnames == ["Anleitung", "EPD", *epd_excel.LISTEN, "Listen"]
    codes = [r[0] for r in wb["Indikatoren"].iter_rows(min_row=3, values_only=True) if r[0]]
    assert codes == [z[0] for z in epd_excel.STANDARD_INDIKATOREN]
    assert wb["EPD"].data_validations.dataValidation, "Auswahllisten im Blatt EPD fehlen"


@pytest.mark.parametrize("serie", FAELLE)
def test_befuellte_excel_ist_gueltig(arbeitsordner, serie):
    _, fehler, _ = epd_excel.pruefe(epd_excel.lese_arbeitsmappe(excel(arbeitsordner, serie)),
                                    basis_ordner=arbeitsordner / serie)
    assert fehler == []


def test_vokabular_der_vorlage_entspricht_template():
    baum, _ = excel_to_aasx.lade_template(TEMPLATE)
    tpl, _ = excel_to_aasx.template_vokabular(baum)
    for name in excel_to_aasx.GESCHLOSSEN:
        assert set(tpl[name]) == set(epd_excel.VOKABULAR[name]), name


@pytest.mark.parametrize("serie", FAELLE)
def test_aasx_werte_und_vokabular(erzeugt, serie):
    _, _, anzahl = FAELLE[serie]
    dateien = sorted((erzeugt / serie).glob("*.aasx"))
    assert len(dateien) == anzahl
    for datei in dateien:
        baum, sm, namen = lies_aasx(datei)
        artikel = kinder(kinder(sm)["manufacturerProduct"])["productArticleNumberOfManufacturer"].findtext(NS + "value")
        werte, status = werte_aus_aasx(sm)
        assert werte == erwartete_werte(excel(erzeugt.parent, serie), artikel), datei.name
        assert {s["valueStatus"] for s in status} <= {"Declared", "NotDeclared"}
        assert all(("value" in s) == (s["valueStatus"] == "Declared") for s in status), "NotDeclared mit Wert"
        stufen = {kinder(x)["stageCode"].findtext(NS + "value"): kinder(x)["stageStatus"].findtext(NS + "value")
                  for x in kinder(kinder(sm)["LCAMethodology"])["DeclaredStages"].find(NS + "value")}
        assert set(stufen.values()) <= {"Assessed", "NotAssessed"}
        assert stufen["C1"] == "NotAssessed" and stufen["C2-C4"] == "Assessed" and stufen["B6"] == "Assessed"
        assert kinder(kinder(sm)["EPDScope"])["scopeType"].findtext(NS + "value") == "Class"
        pv = kinder(kinder(sm)["programOperatorVerification"])
        assert pv["verificationType"].findtext(NS + "value") == "ThirdPartyVerification"
        assert any(n.endswith(".pdf") for n in namen), "EPD-PDF nicht eingebettet"
        assert sm.findtext(NS + "id") == f"https://wago.com/ids/submodel/{artikel}/EPD"


@pytest.mark.parametrize("serie", FAELLE)
def test_aas_core3_konform(erzeugt, serie):
    from aas_core3 import verification, xmlization
    for datei in (erzeugt / serie).glob("*.aasx"):
        with zipfile.ZipFile(datei) as z:
            env = xmlization.environment_from_str(z.read("aasx/data.xml").decode("utf-8"))
        befunde = [f"{f.path}: {f.cause}" for f in verification.verify(env)]
        assert [b for b in befunde if "AASd-120" not in b] == [], datei.name


def _kaputt(arbeitsordner, name, aenderung):
    """Kopie von EPD_221.xlsx mit einer Aenderung (Blatt, Zelle, Wert) bzw. Funktion(wb)."""
    ziel = arbeitsordner / "221" / f"kaputt_{name}.xlsx"
    wb = load_workbook(excel(arbeitsordner, "221"))
    aenderung(wb)
    wb.save(ziel)
    return ziel


def _zeile_mit(ws, spalte, wert):
    return next(r for r in range(3, ws.max_row + 1) if ws.cell(r, spalte).value == wert)


def _epd_zelle(wb, schluessel):
    ws = wb["EPD"]
    return next(ws.cell(r, 4) for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value == schluessel)


FEHLERFAELLE = {
    "scope_vokabular": (lambda wb: setattr(_epd_zelle(wb, "scope_typ"), "value", "Family"),
                        "nicht im Vokabular scopeType"),
    "stage_status": (lambda wb: setattr(wb["Module"].cell(3, 2), "value", "Declared"),
                     "nicht im Vokabular stageStatus"),
    "datum": (lambda wb: setattr(_epd_zelle(wb, "gueltig_bis"), "value", "2020-01-01"),
              "nicht nach dem Ausgabedatum"),
    "zwei_referenzen": (lambda wb: setattr(wb["Produkte"].cell(4, 2), "value", "ja"),
                        "genau ein Referenzprodukt"),
    "wert_kein_zahl": (lambda wb: setattr(wb["Ergebnisse"].cell(3, 4), "value", "abc"), "kein gültiger Wert"),
    "unbekannter_code": (lambda wb: setattr(wb["Ergebnisse"].cell(3, 2), "value", "GWP-xyz"),
                         "nicht im Blatt Indikatoren"),
    "faktor_fehlt": (lambda wb: wb["Faktoren"].delete_rows(3), "Faktoren fehlen"),
    "wert_bei_notassessed": (lambda wb: setattr(wb["Ergebnisse"].cell(3, 3), "value", "B1"),
                             "als NotAssessed markiert"),
    "platzhalter": (lambda wb: setattr(_epd_zelle(wb, "de_beschreibung"), "value", "1 item {artikelnr}"),
                    "unbekannte Platzhalter"),
    "pflichtfeld": (lambda wb: setattr(_epd_zelle(wb, "registriernummer"), "value", None), "Pflichtfeld ist leer"),
}


@pytest.mark.parametrize("fall", FEHLERFAELLE)
def test_fehler_in_excel_werden_erkannt(arbeitsordner, fall, capsys):
    aenderung, erwartet = FEHLERFAELLE[fall]
    datei = _kaputt(arbeitsordner, fall, aenderung)
    out = arbeitsordner / "out_kaputt" / fall
    rc = excel_to_aasx.main([str(datei), "--out", str(out), "--template", str(TEMPLATE)])
    ausgabe = capsys.readouterr().out
    assert rc == 2, ausgabe
    assert erwartet in ausgabe, ausgabe
    assert not list(out.glob("*.aasx")) if out.exists() else True
