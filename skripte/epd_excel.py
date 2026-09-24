#!/usr/bin/env python3
"""Standard-Excel für EPD-Teilmodelle (Vorlage v1.0): Aufbau, Schreiben, Lesen, Prüfen.

Diese Datei ist der "Vertrag" zwischen fachlicher Eingabe (Excel) und dem
generischen Generator excel_to_aasx.py. Sie enthält:

  EPD_FELDER / LISTEN   welche Blätter, Felder und Spalten die Excel hat
  schreibe_arbeitsmappe leere Vorlage oder befuellte Excel im selben Layout
  lese_arbeitsmappe     Excel -> Rohdaten
  pruefe                Rohdaten -> geprüftes Modell + Fehler/Warnungen

Leere Vorlage erzeugen:
    python skripte/epd_excel.py --leer vorlagen/EPD_Vorlage_v1.xlsx

Nur prüfen (ohne AASX zu erzeugen):
    python skripte/epd_excel.py --pruefe epd/221/EPD_221.xlsx

Aendert sich die Vorlage, wird VORLAGEN_VERSION erhöht. Der Generator lehnt
Excel-Dateien mit anderer Hauptversion ab.
"""

import argparse
import datetime as dt
import re
import string
import sys
from pathlib import Path
from typing import NamedTuple

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

VORLAGEN_VERSION = "1.0"

# --------------------------------------------------------------- Vokabular
# Geschlossene Listen stammen aus den ConceptDescriptions (valueList) des
# AASX-Templates epd-type-iii-submodel-template.aasx. excel_to_aasx.py liest
# sie zur Laufzeit erneut aus dem Template und meldet Abweichungen.

VOKABULAR = {
    "scopeType": ["Class", "Instance"],                        # CD EPDScope/Type
    "stageStatus": ["Assessed", "NotAssessed"],                # CD StageStatus
    "valueStatus": ["Declared", "NotDeclared"],                # CD ValueStatus (setzt der Generator)
    "siteType": ["Fixed", "Mobile"],                           # CD ManufacturingSite/Type
    "resultCategory": ["ImpactIndicator", "ResourceUse", "Waste", "OutputFlow"],  # CD ResultCategory
    # Kein CD-Vokabular im Template; Beispielwert des Templates ist
    # "ThirdPartyVerification", das Gegenstueck ist analog gebildet.
    "verificationType": ["ThirdPartyVerification", "InternalVerification"],
    # offene Liste (CD StageCode, "Additional values may be used"), plus
    # gaengige Aggregate und die Summe
    "stageCode": ["A1", "A2", "A3", "A1-A3", "A4", "A5", "B1", "B2", "B3", "B4", "B5", "B6", "B7",
                  "C1", "C2", "C3", "C4", "C2-C4", "C1-C4", "D", "Total"],
    # Hilfslisten der Excel
    "datumsgenauigkeit": ["Tag", "Monat"],
    "werteherkunft": ["Deklariert", "Faktor"],
    "janein": ["ja", "nein"],
    "normart": ["Norm", "PCR", "PSR"],
    "pcrStatus": ["mandatory", "optional"],
}
OFFENE_LISTEN = {"stageCode"}          # andere Werte erlaubt (Warnung statt Fehler)
SUMMEN_STUFE = "Total"                 # nie in DeclaredStages, nur mit --mit-total ausgegeben

# Platzhalter fuer Textfelder, werden je erzeugter AASX ersetzt
PLATZHALTER = {
    "artikel": "Artikelnummer der erzeugten Variante",
    "referenzprodukt": "Artikelnummer des Referenzprodukts",
    "produktname": "Produktname (Blatt EPD bzw. Spalte Produktname)",
    "serie": "Spalte Serie des Produkts",
    "registriernummer": "Registriernummer der EPD",
    "masse_kg": "Masse der deklarierten Einheit in kg",
    "lca_referenzjahr": "Referenzjahr der LCA",
    "herkunft": "Englischer Satzteil zur Herkunft der Werte (deklariert / über Faktoren hochgerechnet)",
    "produkte": "Komma-Liste aller Artikelnummern der Excel",
}


class Feld(NamedTuple):
    schluessel: str
    bereich: str
    label: str
    pflicht: bool
    typ: str                 # text | zahl | ganzzahl | datum | jahr | auswahl | liste | url | email
    auswahl: str = ""        # Schluessel in VOKABULAR
    beispiel: str = ""
    beschreibung: str = ""
    aas: str = ""
    standard: object = None  # Vorbelegung in der leeren Vorlage (WAGO-weite Konstanten)


EPD_FELDER = [
    # -- Verwaltung
    Feld("vorlagen_version", "Vorlage", "Version der Excel-Vorlage", True, "text",
         beispiel=VORLAGEN_VERSION, beschreibung="Nicht ändern. Der Generator prüft die Hauptversion.",
         standard=VORLAGEN_VERSION),
    Feld("id_basis", "Identifikatoren", "Basis-URL der Identifier", True, "url",
         beispiel="https://wago.com/ids",
         beschreibung="Präfix für Teilmodell-, AAS- und Asset-Id: {basis}/submodel/{artikel}/{kennung}, "
                      "{basis}/aas/{artikel}, {basis}/asset/{artikel}",
         aas="Submodel/id, AAS/id, globalAssetId", standard="https://wago.com/ids"),
    Feld("submodel_kennung", "Identifikatoren", "Kennung des Teilmodells in der Id", True, "text",
         beispiel="EPD", aas="Submodel/id", standard="EPD"),
    Feld("dateiname_praefix", "Identifikatoren", "Präfix der AASX-Dateinamen", True, "text",
         beispiel="wago", beschreibung="Ergibt <praefix>-<artikel>-epd-submodel.aasx", standard="wago"),
    Feld("signifikante_stellen", "Identifikatoren", "Signifikante Stellen hochgerechneter Werte", True,
         "ganzzahl", beispiel="3",
         beschreibung="Werte aus Faktoren (Referenzwert x Faktor) werden auf so viele Stellen gerundet "
                      "wie die Werte der Quelle.", standard=3),
    # -- IdentificationPublication
    Feld("registriernummer", "Identifikation & Veröffentlichung", "Registriernummer der EPD", True, "text",
         beispiel="WAGO-00001-V01.01-EN", aas="identificationPublication/documentIds/documentIdentifier"),
    Feld("dokument_domaene", "Identifikation & Veröffentlichung", "Vergabestelle der Registriernummer", True,
         "text", beispiel="PEP ecopassport", aas="identificationPublication/documentIds/documentDomainId"),
    Feld("programm_version", "Identifikation & Veröffentlichung", "Version beim Programmbetreiber", True,
         "text", beispiel="01.01", aas="identificationPublication/programOperatorVersion"),
    Feld("ausgabedatum", "Identifikation & Veröffentlichung", "Ausgabedatum", True, "datum",
         beispiel="2026-02-01", beschreibung="JJJJ-MM-TT. Nennt die EPD nur Monat/Jahr: Tag 01 eintragen "
                                             "und Datumsgenauigkeit = Monat.",
         aas="identificationPublication/dateOfIssue"),
    Feld("gueltig_bis", "Identifikation & Veröffentlichung", "Gültig bis", True, "datum",
         beispiel="2031-02-01", aas="identificationPublication/validUntil"),
    Feld("datumsgenauigkeit", "Identifikation & Veröffentlichung", "Genauigkeit der Datumsangaben", True,
         "auswahl", "datumsgenauigkeit", beispiel="Monat",
         beschreibung="Monat = Tag ist ergänzt; erzeugt Qualifier EPD/DatePrecision"),
    Feld("sprachen", "Identifikation & Veröffentlichung", "Sprachen der EPD", True, "liste",
         beispiel="en", beschreibung="ISO 639-1, mehrere mit ; trennen", aas="identificationPublication/languages"),
    Feld("quelldatenformat", "Identifikation & Veröffentlichung", "Format der Quelldaten", False, "text",
         beispiel="PEP ecopassport PDF", aas="identificationPublication/sourceDataFormat"),
    Feld("deklarations_url", "Identifikation & Veröffentlichung", "Öffentliche URL der EPD", False, "url",
         beispiel="https://register.pep-ecopassport.org/...", aas="identificationPublication/declarationUrl"),
    Feld("epd_dokument", "Identifikation & Veröffentlichung", "EPD-PDF (Dateiname)", False, "text",
         beispiel="WAGO-00001-V01.01-EN.pdf",
         beschreibung="Pfad relativ zur Excel. Wird in die AASX eingebettet. Im Template Pflicht (epdDocument).",
         aas="epdDocument"),
    # -- ProgramOperatorVerification
    Feld("programmbetreiber", "Programmbetreiber & Verifizierung", "Programmbetreiber", True, "text",
         beispiel="Association P.E.P. (PEP ecopassport)", aas="programOperatorVerification/programOperator/Company"),
    Feld("programmbetreiber_email", "Programmbetreiber & Verifizierung", "E-Mail Programmbetreiber", False,
         "email", beispiel="contact@pep-ecopassport.org",
         aas="programOperatorVerification/programOperator/EmailAddress"),
    Feld("verifizierungsart", "Programmbetreiber & Verifizierung", "Art der Verifizierung", True, "auswahl",
         "verificationType", beispiel="ThirdPartyVerification",
         aas="programOperatorVerification/verificationType"),
    Feld("verifizierer_akkreditierung", "Programmbetreiber & Verifizierung", "Akkreditierungsnr. Verifizierer",
         False, "text", beispiel="VH44",
         aas="programOperatorVerification/thirdPartyVerifier/verifierAccreditationId"),
    Feld("verifizierungsdatum", "Programmbetreiber & Verifizierung", "Datum der Verifizierung", False, "datum",
         beispiel="2026-01-15", aas="programOperatorVerification/verificationDate"),
    Feld("epd_ersteller", "Programmbetreiber & Verifizierung", "Ersteller der EPD / Ökobilanz", False, "text",
         beispiel="WAGO Kontakttechnik GmbH & Co. KG", aas="programOperatorVerification/epdDeveloper/Company"),
    # -- ManufacturerProduct
    Feld("hersteller", "Hersteller & Produkt", "Hersteller", True, "text",
         beispiel="WAGO Kontakttechnik GmbH & Co. KG", aas="manufacturerProduct/manufacturerName",
         standard="WAGO Kontakttechnik GmbH & Co. KG"),
    Feld("produktname", "Hersteller & Produkt", "Produktname (Familie)", True, "text",
         beispiel="221 Series / Splicing Connector with Levers",
         beschreibung="Je Produkt überschreibbar (Blatt Produkte, Spalte Produktname)",
         aas="manufacturerProduct/productName"),
    Feld("referenzlebensdauer_jahre", "Hersteller & Produkt", "Referenzlebensdauer (Jahre)", False, "zahl",
         beispiel="30", aas="manufacturerProduct/referenceServiceLife"),
    Feld("geltungsbereich", "Hersteller & Produkt", "Geltungsbereich (Länder/Regionen)", False, "liste",
         beispiel="001", beschreibung="ISO 3166-1 (DE) oder UN M49 (001 = Welt), mehrere mit ; trennen",
         aas="manufacturerProduct/applicableJurisdictions"),
    Feld("herstellungsbeschreibung", "Hersteller & Produkt", "Beschreibung der Herstellung", False, "text",
         beispiel="Production at WAGO facilities in Germany (injection moulding, metal forming, assembly).",
         aas="manufacturerProduct/manufacturingDescription"),
    # -- EPDScope
    Feld("scope_typ", "Geltungsbereich der EPD", "Scope-Typ", True, "auswahl", "scopeType", beispiel="Class",
         beschreibung="Class = EPD für eine Produktfamilie (mehrere Artikel), Instance = EPD für genau "
                      "einen Artikel", aas="EPDScope/scopeType"),
    Feld("scope_beschreibung", "Geltungsbereich der EPD", "Beschreibung des Geltungsbereichs", False, "text",
         beispiel="Product family PEP {registriernummer}, reference product {referenzprodukt}. "
                  "This submodel describes {artikel}: {herkunft}.",
         beschreibung="Platzhalter erlaubt, siehe Anleitung", aas="EPDScope/scopeDescription"),
    # -- LCAMethodology
    Feld("lca_referenzjahr", "LCA-Methodik", "Referenzjahr der Daten", True, "jahr", beispiel="2024",
         aas="LCAMethodology/lcaReferenceYear"),
    Feld("zeitliche_repraesentativitaet", "LCA-Methodik", "Zeitliche Repräsentativität", True, "text",
         beispiel="Primary production data refer to calendar year {lca_referenzjahr}.",
         aas="LCAMethodology/timeRepresentativenessDescription"),
    Feld("geografischer_scope", "LCA-Methodik", "Geografischer Geltungsbereich der LCA", True, "text",
         beispiel="GLO", beschreibung="Code, z. B. GLO, RER, DE", aas="LCAMethodology/geographicalScope"),
    Feld("geografische_beschreibung", "LCA-Methodik", "Geografische Repräsentativität", True, "text",
         beispiel="Primary data from WAGO facilities in Germany; background data ecoinvent (GLO/RER).",
         aas="LCAMethodology/geographicalDescription"),
    Feld("technologie_beschreibung", "LCA-Methodik", "Technologische Repräsentativität", True, "text",
         beispiel="Actual manufacturing technologies (injection moulding, metal forming, assembly).",
         aas="LCAMethodology/technologyDescription"),
    Feld("energiemodell", "LCA-Methodik", "Strommix der Herstellung", False, "text", beispiel="DE",
         beschreibung="Code des Strommixes, z. B. DE, EU", aas="LCAMethodology/lcaEnergyModel"),
    # -- FunctionalUnit / DeclaredUnit
    Feld("fe_menge", "Funktionelle Einheit", "Menge", False, "zahl", beispiel="1", aas="FunctionalUnit/quantity"),
    Feld("fe_einheit", "Funktionelle Einheit", "Einheit", False, "text", beispiel="A", aas="FunctionalUnit/unit"),
    Feld("fe_beschreibung", "Funktionelle Einheit", "Beschreibung", False, "text",
         beispiel="1 A over 30 years at a use rate of 70 %", aas="FunctionalUnit/description"),
    Feld("de_menge", "Deklarierte Einheit", "Menge", False, "zahl", beispiel="1", aas="DeclaredUnit/quantity"),
    Feld("de_einheit", "Deklarierte Einheit", "Einheit", False, "text", beispiel="item", aas="DeclaredUnit/unit"),
    Feld("de_beschreibung", "Deklarierte Einheit", "Beschreibung", False, "text",
         beispiel="1 item {artikel} including packaging; mass {masse_kg} kg per declared unit.",
         aas="DeclaredUnit/description"),
    # -- ContentDeclaration
    Feld("recyclinganteil_prozent", "Inhaltsdeklaration", "Recyclinganteil (%)", False, "zahl", beispiel="12.5",
         aas="ContentDeclaration/recycledContentPercentage"),
]
FELD = {f.schluessel: f for f in EPD_FELDER}


class Spalte(NamedTuple):
    schluessel: str
    kopf: str
    pflicht: bool
    typ: str
    auswahl: str = ""
    hinweis: str = ""        # steht in Zeile 2 des Blatts
    breite: int = 16


LISTEN = {
    "Produkte": [
        Spalte("artikel", "Artikelnummer", True, "text", hinweis="z. B. 221-422", breite=14),
        Spalte("referenz", "Referenzprodukt", True, "auswahl", "janein",
               hinweis="genau ein Produkt = ja", breite=13),
        Spalte("herkunft", "Werteherkunft", True, "auswahl", "werteherkunft",
               hinweis="Deklariert = Werte im Blatt Ergebnisse; Faktor = Referenz x Blatt Faktoren", breite=14),
        Spalte("serie", "Serie", False, "text", hinweis="Freitext, Platzhalter {serie}", breite=30),
        Spalte("produktname", "Produktname", False, "text",
               hinweis="leer = Produktname aus Blatt EPD", breite=24),
        Spalte("beschreibung", "Produktbeschreibung", False, "text",
               hinweis="-> productDescription", breite=45),
        Spalte("masse_kg", "Masse deklarierte Einheit (kg)", True, "zahl",
               hinweis="inkl. Verpackung, z. B. 0.001958", breite=16),
        Spalte("produktmasse_kg", "Produktmasse ohne Verpackung (kg)", False, "zahl",
               hinweis="-> ContentDeclaration/productMass", breite=16),
        Spalte("biogen_produkt", "Biogener Kohlenstoff Produkt (kg C)", False, "zahl", breite=16),
        Spalte("biogen_verpackung", "Biogener Kohlenstoff Verpackung (kg C)", False, "zahl", breite=16),
        Spalte("bild", "Produktbild (Datei)", False, "text", hinweis="jpg/png, relativ zur Excel", breite=26),
        Spalte("erzeugen", "AASX erzeugen", False, "auswahl", "janein", hinweis="leer = ja", breite=10),
    ],
    "Module": [
        Spalte("stageCode", "stageCode", True, "auswahl", "stageCode",
               hinweis="Reihenfolge = Reihenfolge in der AASX", breite=12),
        Spalte("stageStatus", "stageStatus", True, "auswahl", "stageStatus",
               hinweis="Assessed | NotAssessed", breite=14),
        Spalte("hinweis", "Hinweis", False, "text",
               hinweis="z. B. Zuordnung der Quellspalte; -> Qualifier EPD/StageNote", breite=90),
    ],
    "Indikatoren": [
        Spalte("code", "indicatorCode", True, "text", hinweis="Vokabular des Templates, z. B. GWP-total", breite=15),
        Spalte("kategorie", "resultCategory", True, "auswahl", "resultCategory", breite=16),
        Spalte("name", "indicatorName", True, "text", breite=48),
        Spalte("einheit", "unit", True, "text", hinweis="Registry-Schreibweise, z. B. kg CO2 eq.", breite=16),
        Spalte("basiseinheit", "baseUnit", False, "text", hinweis="nur bei abgeleiteten Einheiten", breite=10),
        Spalte("status", "Status laut PCR", False, "auswahl", "pcrStatus",
               hinweis="-> Qualifier EPD/IndicatorStatus", breite=12),
        Spalte("quelle", "Bezeichnung in der Quelle", False, "text",
               hinweis="Name wie in der EPD gedruckt (Zuordnungshilfe)", breite=70),
    ],
    "Ergebnisse": [
        Spalte("artikel", "Artikelnummer", True, "text", hinweis="Produkt mit Werteherkunft Deklariert", breite=14),
        Spalte("code", "indicatorCode", True, "text", hinweis="aus Blatt Indikatoren", breite=15),
        Spalte("stage", "stageCode", True, "text", hinweis="aus Blatt Module oder Total", breite=10),
        Spalte("wert", "Wert", True, "zahl", hinweis="in der Einheit des Indikators; 0 ist ein Wert!", breite=14),
        Spalte("einheit_quelle", "Einheit in der Quelle", False, "text",
               hinweis="wie gedruckt, falls abweichend -> Qualifier", breite=18),
    ],
    "Faktoren": [
        Spalte("artikel", "Artikelnummer", True, "text", hinweis="Produkt mit Werteherkunft Faktor", breite=14),
        Spalte("code", "indicatorCode", True, "text", breite=15),
        Spalte("stage", "stageCode", True, "text", breite=10),
        Spalte("faktor", "Faktor", True, "zahl", hinweis="Wert = Referenzwert x Faktor", breite=10),
    ],
    "Standorte": [
        Spalte("name", "siteName", True, "text", breite=34),
        Spalte("typ", "siteType", True, "auswahl", "siteType", breite=10),
        Spalte("betreiber", "siteOperator", False, "text", breite=30),
        Spalte("land", "countryCode", False, "text", hinweis="ISO 3166-1 alpha-2", breite=11),
        Spalte("region", "region", False, "text", breite=18),
        Spalte("ort", "locality", False, "text", breite=16),
        Spalte("breite", "latitude", False, "zahl", breite=10),
        Spalte("laenge", "longitude", False, "zahl", breite=10),
        Spalte("id_schema", "identifierScheme", False, "text", hinweis="z. B. Internal", breite=14),
        Spalte("id_wert", "identifierValue", False, "text", breite=16),
    ],
    "Normen_PCR": [
        Spalte("art", "Art", True, "auswahl", "normart", hinweis="Norm -> standardsCompliance; PCR/PSR -> "
                                                                  "productCategoryRules", breite=8),
        Spalte("bezeichnung", "Bezeichnung", True, "text", hinweis="z. B. ISO 14025:2006", breite=40),
    ],
}
BLATT_BESCHREIBUNG = {
    "EPD": "Kopfdaten der EPD (ein Wert je Feld)",
    "Produkte": "Alle Artikel der EPD; je Zeile entsteht eine AASX",
    "Module": "Lebenszyklusmodule und ihr Status (DeclaredStages)",
    "Indikatoren": "Indikatorkatalog: Code, Kategorie, Name, Einheit",
    "Ergebnisse": "Werte je Artikel, Indikator und Modul (Langformat)",
    "Faktoren": "Multiplikationsfaktoren für Varianten (optional)",
    "Standorte": "Fertigungsstandorte (optional)",
    "Normen_PCR": "Normen, PCR und PSR",
    "Listen": "Auswahllisten (nicht ändern)",
}

# Vorbelegung der leeren Vorlage: EN 15804+A2-Indikatoren in der Schreibweise
# des Template-Vokabulars bzw. der Indikator-Registry. "quelle" = Bezeichnung
# im PEP ecopassport. Zusaetzliche Indikatoren duerfen ergaenzt werden.
STANDARD_INDIKATOREN = [
    # code, kategorie, name, einheit, basiseinheit, status, quelle
    ("GWP-total", "ImpactIndicator", "Global Warming Potential (total)", "kg CO2 eq.", "kg", "mandatory",
     "Climate change - total"),
    ("GWP-fossil", "ImpactIndicator", "Global Warming Potential (fossil)", "kg CO2 eq.", "kg", "mandatory",
     "Climate change - fossil fuels"),
    ("GWP-biogenic", "ImpactIndicator", "Global Warming Potential (biogenic)", "kg CO2 eq.", "kg", "mandatory",
     "Climate change - biogenics"),
    ("GWP-luluc", "ImpactIndicator", "Global Warming Potential (land use and land use change)", "kg CO2 eq.",
     "kg", "mandatory", "Climate change - land use and land use transformation"),
    ("ODP", "ImpactIndicator", "Ozone Depletion Potential", "kg CFC-11 eq.", "kg", "mandatory", "Ozone depletion"),
    ("AP", "ImpactIndicator", "Acidification Potential", "mol H+ eq.", "mol", "mandatory", "Acidification (AP)"),
    ("EP-freshwater", "ImpactIndicator", "Eutrophication Potential (freshwater)", "kg P eq.", "kg", "mandatory",
     "Freshwater eutrophication"),
    ("EP-marine", "ImpactIndicator", "Eutrophication Potential (marine)", "kg N eq.", "kg", "mandatory",
     "Marine aquatic eutrophication"),
    ("EP-terrestrial", "ImpactIndicator", "Eutrophication Potential (terrestrial)", "mol N eq.", "mol",
     "mandatory", "Terrestrial eutrophication"),
    ("POCP", "ImpactIndicator", "Photochemical Ozone Creation Potential", "kg NMVOC eq.", "kg", "mandatory",
     "Photochemical ozone formation"),
    ("ADPE", "ImpactIndicator", "Abiotic Depletion Potential (elements)", "kg Sb eq.", "kg", "mandatory",
     "Abiotic resource depletion - elements or resource depletion - metals and minerals"),
    ("ADPF", "ImpactIndicator", "Abiotic Depletion Potential (fossil fuels)", "MJ", "", "mandatory",
     "Abiotic resource depletion - fossil fuels or resource depletion - fossils"),
    ("WDP", "ImpactIndicator", "Water Deprivation Potential", "m3 eq.", "m3", "mandatory", "Water requirement"),
    ("Total-PE", "ImpactIndicator", "Total Primary Energy", "MJ", "", "optional",
     "Total use of primary energy during the life cycle"),
    ("PM", "ImpactIndicator", "Particulate Matter Formation", "disease incidence", "", "optional",
     "Emission of fine particles"),
    ("IR", "ImpactIndicator", "Ionising Radiation", "kBq U-235 eq.", "kBq", "optional",
     "Ionizing radiation, human health"),
    ("ETP-fw", "ImpactIndicator", "Ecotoxicity (freshwater)", "CTUe", "", "optional", "Ecotoxicity (fresh water)"),
    ("HTP-c", "ImpactIndicator", "Human Toxicity Potential (cancer)", "CTUh", "", "optional",
     "Human toxicity, carcinogenic effects"),
    ("HTP-nc", "ImpactIndicator", "Human Toxicity Potential (non-cancer)", "CTUh", "", "optional",
     "Human toxicity, non-carcinogenic effects"),
    ("SQP", "ImpactIndicator", "Soil Quality Potential", "dimensionless", "", "optional",
     "Impacts related to land use/soil quality"),
    ("PERE", "ResourceUse", "Renewable Primary Energy (excluding raw materials)", "MJ", "", "mandatory",
     "Use of renewable primary energy, excluding renewable primary energy resources used as raw materials"),
    ("PERM", "ResourceUse", "Renewable Primary Energy used as raw materials", "MJ", "", "mandatory",
     "Use of renewable primary energy resources used as raw materials"),
    ("PERT", "ResourceUse", "Total Renewable Primary Energy", "MJ", "", "mandatory",
     "Total use of renewable primary energy resources"),
    ("PENRE", "ResourceUse", "Non-renewable Primary Energy (excluding raw materials)", "MJ", "", "mandatory",
     "Use of non-renewable primary energy, excluding non-renewable primary energy resources used as raw materials"),
    ("PENRM", "ResourceUse", "Non-renewable Primary Energy used as raw materials", "MJ", "", "mandatory",
     "Use of non-renewable primary energy resources used as raw materials"),
    ("PENRT", "ResourceUse", "Total Non-renewable Primary Energy", "MJ", "", "mandatory",
     "Total use of non-renewable primary energy resources"),
    ("SM", "ResourceUse", "Use of Secondary Materials", "kg", "", "mandatory", "Use of secondary materials"),
    ("RSF", "ResourceUse", "Use of Renewable Secondary Fuels", "MJ", "", "mandatory",
     "Use of renewable secondary fuels"),
    ("NRSF", "ResourceUse", "Use of Non-renewable Secondary Fuels", "MJ", "", "mandatory",
     "Use of non-renewable secondary fuels"),
    ("FW", "ResourceUse", "Net Use of Fresh Water", "m3", "", "mandatory", "Net use of fresh water"),
    ("HWD", "Waste", "Hazardous Waste Disposed", "kg", "", "mandatory", "Hazardous waste disposed of"),
    ("NHWD", "Waste", "Non-Hazardous Waste Disposed", "kg", "", "mandatory", "Non-hazardous waste disposed of"),
    ("RWD", "Waste", "Radioactive Waste Disposed", "kg", "", "mandatory", "Radioactive waste disposed of"),
    ("CRU", "OutputFlow", "Components for Re-use", "kg", "", "mandatory", "Components for reuse"),
    ("MFR", "OutputFlow", "Materials for Recycling", "kg", "", "mandatory", "Materials for recycling"),
    ("MER", "OutputFlow", "Materials for Energy Recovery", "kg", "", "mandatory", "Materials for energy recovery"),
    ("EE", "OutputFlow", "Exported Energy", "MJ", "", "mandatory", "Exported energy"),
]
# Module nach EN 15804, Status bewusst leer: muss je EPD entschieden werden
STANDARD_MODULE = ["A1-A3", "A4", "A5", "B1", "B2", "B3", "B4", "B5", "B6", "B7", "C1", "C2", "C3", "C4", "D"]


# --------------------------------------------------------------- Schreiben

SCHRIFT = "Arial"
F_KOPF = PatternFill("solid", fgColor="1F3864")
F_BEREICH = PatternFill("solid", fgColor="D9E1F2")
F_EINGABE = PatternFill("solid", fgColor="FFF2CC")
F_HINWEIS = PatternFill("solid", fgColor="F2F2F2")
RAHMEN = Border(bottom=Side(style="thin", color="BFBFBF"))
DATEN_AB = 3                            # Listenblaetter: Zeile 1 Kopf, Zeile 2 Hinweis
LEERZEILEN = 300                        # so weit reichen Auswahllisten in leeren Blaettern


def _font(**kw):
    return Font(name=SCHRIFT, **{"size": 10, **kw})


def _listen_bereiche(ws_listen):
    """Listen-Blatt schreiben, liefert {vokabular: '=Listen!$A$2:$A$n'}."""
    bereiche = {}
    for i, (name, werte) in enumerate(VOKABULAR.items(), start=1):
        sp = get_column_letter(i)
        ws_listen.cell(1, i, name).font = _font(bold=True, color="FFFFFF")
        ws_listen.cell(1, i).fill = F_KOPF
        for j, w in enumerate(werte, start=2):
            ws_listen.cell(j, i, w).font = _font()
        ws_listen.column_dimensions[sp].width = max(14, max(len(w) for w in werte) + 2)
        bereiche[name] = f"=Listen!${sp}$2:${sp}${len(werte) + 1}"
    ws_listen.cell(len(max(VOKABULAR.values(), key=len)) + 3, 1,
                   "Quelle: ConceptDescriptions (valueList) des AASX-Templates. Nicht ändern.").font = \
        _font(italic=True, color="808080")
    return bereiche


def _validierung(ws, bereich, formel, offen=False, titel=""):
    dv = DataValidation(type="list", formula1=formel, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorStyle = "warning" if offen else "stop"
    dv.errorTitle = "Wert nicht in der Liste"
    dv.error = ("Wert steht nicht in der Liste. Zusätzliche Werte sind erlaubt, werden aber gemeldet."
                if offen else "Bitte einen Wert aus der Liste wählen (Template-Vokabular).")
    if titel:
        dv.promptTitle, dv.prompt, dv.showInputMessage = titel[:32], "Auswahl aus Liste", True
    ws.add_data_validation(dv)
    dv.add(bereich)


def _anleitung(ws):
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 120
    zeilen = [
        ("titel", f"EPD -> AASX: Standard-Excel (Vorlage v{VORLAGEN_VERSION})"),
        ("text", "Eine Excel je EPD. Daraus erzeugt excel_to_aasx.py für jeden Artikel im Blatt Produkte eine "
                 "AASX mit dem Teilmodell EPD Type III."),
        ("text", "Vollständige Beispiele: epd/221/EPD_221.xlsx und epd/2273/EPD_2273.xlsx (aus den WAGO-PEPs erzeugt)."),
        ("text", "Ablage: je EPD ein Ordner epd/<serie>/ mit PDF, dieser Excel, produktbilder/ und den erzeugten "
                 "AASX in aasx/. Dateinamen in der Excel (PDF, Bilder) sind relativ zu diesem Ordner."),
        ("", ""),
        ("abschnitt", "So wird ausgefüllt"),
        ("schritt", "1. Blatt EPD: gelbe Zellen in Spalte Wert füllen. Pflichtfelder: Spalte Pflicht = ja."),
        ("schritt", "2. Blatt Produkte: je Artikel eine Zeile. Genau ein Referenzprodukt (ja)."),
        ("schritt", "3. Blatt Module: Status je Lebenszyklusmodul festlegen (Assessed / NotAssessed). "
                    "Aggregierte Module (z. B. A1-A3, C2-C4) als eigene Zeile."),
        ("schritt", "4. Blatt Indikatoren: vorbelegter Katalog (EN 15804+A2). Nur ergänzen, wenn die EPD "
                    "weitere Indikatoren enthält."),
        ("schritt", "5. Blatt Ergebnisse: je Wert eine Zeile (Artikelnummer, indicatorCode, stageCode, Wert). "
                    "Fehlender Wert = keine Zeile (wird NotDeclared), nicht 0 eintragen."),
        ("schritt", "6. Blatt Faktoren: nur für Produktfamilien. Wert der Variante = Wert des Referenzprodukts x Faktor."),
        ("schritt", "7. Blätter Standorte und Normen_PCR füllen."),
        ("schritt", "8. Prüfen:   python skripte/epd_excel.py --pruefe epd/<serie>/EPD_<serie>.xlsx"),
        ("schritt", "9. Erzeugen:  python skripte/excel_to_aasx.py epd/<serie>/EPD_<serie>.xlsx   (AASX in epd/<serie>/aasx/)"),
        ("", ""),
        ("abschnitt", "Legende"),
        ("legende", "Gelb hinterlegt = Eingabezelle"),
        ("legende", "Spaltenkopf mit * = Pflichtspalte"),
        ("legende", "Zeile 2 der Listenblätter (grau) = Hinweis, nicht überschreiben. Daten ab Zeile 3."),
        ("legende", "Auswahllisten stammen aus dem Vokabular des AASX-Templates (Blatt Listen)."),
        ("", ""),
        ("abschnitt", "Vokabular (Template)"),
        ("text", "stageStatus: Assessed = Modul ist bewertet, NotAssessed = nicht bewertet / nicht anwendbar."),
        ("text", "valueStatus (setzt das Skript): Declared = Wert vorhanden (auch 0), NotDeclared = kein Wert."),
        ("text", "scopeType: Class = EPD für eine Produktfamilie, Instance = EPD für genau einen Artikel."),
        ("text", "verificationType: ThirdPartyVerification (extern) oder InternalVerification."),
        ("", ""),
        ("abschnitt", "Platzhalter in Textfeldern"),
    ] + [("platzhalter", f"{{{k}}}  -  {v}") for k, v in PLATZHALTER.items()] + [
        ("", ""),
        ("abschnitt", "Blätter"),
    ] + [("platzhalter", f"{k}  -  {v}") for k, v in BLATT_BESCHREIBUNG.items()]
    for i, (art, text) in enumerate(zeilen, start=1):
        c = ws.cell(i, 2, text)
        if art == "titel":
            c.font = _font(size=14, bold=True, color="1F3864")
        elif art == "abschnitt":
            c.font = _font(size=11, bold=True, color="1F3864")
        else:
            c.font = _font()
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if art == "legende" and text.startswith("Gelb"):
            ws.cell(i, 1).fill = F_EINGABE
    ws.sheet_view.showGridLines = False


def _epd_blatt(ws, werte, bereiche):
    kopf = ["Bereich", "Feld", "Schlüssel", "Wert", "Pflicht", "Format / Auswahl", "Beispiel", "Beschreibung",
            "AAS-Element"]
    breiten = [26, 36, 24, 60, 8, 22, 40, 60, 48]
    for i, (k, b) in enumerate(zip(kopf, breiten), start=1):
        c = ws.cell(1, i, k)
        c.font, c.fill = _font(bold=True, color="FFFFFF"), F_KOPF
        ws.column_dimensions[get_column_letter(i)].width = b
    ws.freeze_panes = "D2"
    z, bereich_alt = 2, None
    for f in EPD_FELDER:
        if f.bereich != bereich_alt:
            ws.cell(z, 1, f.bereich).font = _font(bold=True)
            for sp in range(1, len(kopf) + 1):
                ws.cell(z, sp).fill = F_BEREICH
            bereich_alt, z = f.bereich, z + 1
        wert = werte.get(f.schluessel, f.standard) if werte is not None else f.standard
        format_text = {"datum": "JJJJ-MM-TT", "jahr": "JJJJ", "zahl": "Zahl", "ganzzahl": "ganze Zahl",
                       "liste": "Werte mit ; trennen", "url": "URL", "email": "E-Mail", "text": "Text"}.get(f.typ, "")
        if f.typ == "auswahl":
            format_text = " | ".join(VOKABULAR[f.auswahl])
        zeile = [f.bereich, f.label, f.schluessel, wert, "ja" if f.pflicht else "", format_text, f.beispiel,
                 f.beschreibung, f.aas]
        for sp, v in enumerate(zeile, start=1):
            c = ws.cell(z, sp, v)
            c.font = _font(color="808080", size=9) if sp in (1, 3, 9) else _font(bold=(sp == 2 and f.pflicht))
            c.alignment = Alignment(wrap_text=sp in (4, 7, 8), vertical="top", horizontal="left" if sp == 4 else None)
            c.border = RAHMEN
        wz = ws.cell(z, 4)
        wz.fill = F_EINGABE
        if f.typ == "datum":
            wz.number_format = "yyyy-mm-dd"
        elif f.typ in ("text", "liste", "url", "email") :
            wz.number_format = "@"
        if f.typ == "auswahl":
            _validierung(ws, f"D{z}", bereiche[f.auswahl], f.auswahl in OFFENE_LISTEN, f.label)
        z += 1


def _listen_blatt(ws, blatt, zeilen, bereiche):
    spalten = LISTEN[blatt]
    for i, s in enumerate(spalten, start=1):
        c = ws.cell(1, i, s.kopf + (" *" if s.pflicht else ""))
        c.font, c.fill = _font(bold=True, color="FFFFFF"), F_KOPF
        c.alignment = Alignment(wrap_text=True, vertical="center")
        h = ws.cell(2, i, s.hinweis or ("Pflicht" if s.pflicht else "optional"))
        h.font, h.fill = _font(italic=True, color="808080", size=9), F_HINWEIS
        h.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = s.breite
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 36
    ws.freeze_panes = f"A{DATEN_AB}"
    for r, zeile in enumerate(zeilen or [], start=DATEN_AB):
        for i, s in enumerate(spalten, start=1):
            v = zeile.get(s.schluessel)
            c = ws.cell(r, i, v)
            c.font = _font()
            if s.typ == "zahl" and blatt in ("Ergebnisse",):
                c.number_format = "0.00E+00"
            elif s.typ == "text":
                c.number_format = "@"
    ende = max(DATEN_AB + len(zeilen or []) + 50, LEERZEILEN)
    for i, s in enumerate(spalten, start=1):
        sp = get_column_letter(i)
        rng = f"{sp}{DATEN_AB}:{sp}{ende}"
        for r in range(DATEN_AB, ende + 1):
            ws.cell(r, i).fill = F_EINGABE if r < DATEN_AB + len(zeilen or []) + 20 else PatternFill()
        if s.typ == "auswahl":
            _validierung(ws, rng, bereiche[s.auswahl], s.auswahl in OFFENE_LISTEN, s.kopf)
    # Querverweise: Artikel, Indikatoren und Module aus den anderen Blaettern
    verweise = {("Ergebnisse", "artikel"): "=Produkte!$A$3:$A$1000",
                ("Faktoren", "artikel"): "=Produkte!$A$3:$A$1000",
                ("Ergebnisse", "code"): "=Indikatoren!$A$3:$A$500",
                ("Faktoren", "code"): "=Indikatoren!$A$3:$A$500",
                ("Ergebnisse", "stage"): bereiche["stageCode"],
                ("Faktoren", "stage"): bereiche["stageCode"]}
    for i, s in enumerate(spalten, start=1):
        formel = verweise.get((blatt, s.schluessel))
        if formel:
            sp = get_column_letter(i)
            _validierung(ws, f"{sp}{DATEN_AB}:{sp}{ende}", formel, offen=True)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(spalten))}{max(DATEN_AB, DATEN_AB + len(zeilen or []) - 1)}"


def schreibe_arbeitsmappe(pfad, epd=None, listen=None):
    """Excel im Vorlagenlayout schreiben.

    epd:    {schluessel: wert} für Blatt EPD (None = leere Vorlage mit Vorbelegung)
    listen: {blatt: [ {spalten_schluessel: wert} ]}; fehlende Blätter werden in der
            leeren Vorlage mit Indikatorkatalog und Modulliste vorbelegt.
    """
    listen = dict(listen or {})
    if epd is None:
        listen.setdefault("Indikatoren", [dict(zip(("code", "kategorie", "name", "einheit", "basiseinheit",
                                                        "status", "quelle"), z)) for z in STANDARD_INDIKATOREN])
        listen.setdefault("Module", [{"stageCode": m} for m in STANDARD_MODULE])
    wb = Workbook()
    ws = wb.active
    ws.title = "Anleitung"
    _anleitung(ws)
    ws_epd = wb.create_sheet("EPD")
    blaetter = {b: wb.create_sheet(b) for b in LISTEN}
    ws_listen = wb.create_sheet("Listen")
    bereiche = _listen_bereiche(ws_listen)
    _epd_blatt(ws_epd, epd, bereiche)
    for b, w in blaetter.items():
        _listen_blatt(w, b, listen.get(b), bereiche)
    ws_listen.sheet_state = "visible"
    wb.properties.title = f"EPD-Vorlage v{VORLAGEN_VERSION}"
    wb.properties.creator = "epd_excel.py"
    wb.save(pfad)


# --------------------------------------------------------------- Lesen

def _kopf_normal(text):
    """Spaltenkopf vergleichbar machen: Pflicht-Stern weg, Kleinschreibung, Umlaute ausgeschrieben."""
    t = re.sub(r"\s*\*\s*$", "", str(text or "")).strip().lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        t = t.replace(a, b)
    return t


def lese_arbeitsmappe(pfad):
    """Excel -> {'epd': {schluessel: (wert, zelle)}, 'listen': {blatt: [ {sp: wert, '_zeile': n} ]},
    'fehler': [strukturfehler]}"""
    wb = load_workbook(pfad, data_only=True)
    fehler = []
    epd = {}
    if "EPD" not in wb.sheetnames:
        fehler.append("Blatt 'EPD' fehlt - ist das eine Excel nach Vorlage v1?")
    else:
        ws = wb["EPD"]
        kopf = [_kopf_normal(c.value) for c in ws[1]]
        try:
            sp_schl, sp_wert = kopf.index("schluessel"), kopf.index("wert")
        except ValueError:
            fehler.append("Blatt EPD: Spalten 'Schlüssel' und 'Wert' nicht gefunden")
        else:
            for zeile in ws.iter_rows(min_row=2):
                schl = zeile[sp_schl].value
                if schl:
                    epd[str(schl).strip()] = (zeile[sp_wert].value, f"EPD!{zeile[sp_wert].coordinate}")
    listen = {}
    for blatt, spalten in LISTEN.items():
        if blatt not in wb.sheetnames:
            fehler.append(f"Blatt '{blatt}' fehlt")
            listen[blatt] = []
            continue
        ws = wb[blatt]
        kopf = {_kopf_normal(c.value): i for i, c in enumerate(ws[1]) if c.value}
        index = {}
        for s in spalten:
            if _kopf_normal(s.kopf) in kopf:
                index[s.schluessel] = kopf[_kopf_normal(s.kopf)]
            elif s.pflicht:
                fehler.append(f"Blatt {blatt}: Pflichtspalte '{s.kopf}' fehlt")
        daten = []
        for zeile in ws.iter_rows(min_row=DATEN_AB):
            werte = {k: zeile[i].value if i < len(zeile) else None for k, i in index.items()}
            if all(v is None or str(v).strip() == "" for v in werte.values()):
                continue
            werte["_zeile"] = zeile[0].row
            werte["_zellen"] = {k: f"{blatt}!{get_column_letter(i + 1)}{zeile[0].row}" for k, i in index.items()}
            daten.append(werte)
        listen[blatt] = daten
    return {"epd": epd, "listen": listen, "fehler": fehler}


# --------------------------------------------------------------- Pruefen

def _leer(v):
    return v is None or (isinstance(v, str) and v.strip() == "") or (isinstance(v, float) and v != v)


def _zahl(v):
    if isinstance(v, bool):
        raise ValueError
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(" ", "")
    if re.fullmatch(r"-?\d+,\d+(?:[Ee][+-]?\d+)?", t):      # Dezimalkomma
        t = t.replace(",", ".")
    return float(t)


def _datum(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    t = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    raise ValueError


def platzhalter_in(text):
    return {f for _, f, _, _ in string.Formatter().parse(str(text)) if f is not None}


def _pruefe_wert(typ, auswahl, v, ort, fehler, warnungen, vokabular):
    """Wert typgerecht umwandeln; Fehler/Warnungen anhaengen. Rückgabe: Wert oder None."""
    if _leer(v):
        return None
    try:
        if typ == "zahl":
            return _zahl(v)
        if typ == "ganzzahl":
            z = _zahl(v)
            if z != int(z):
                raise ValueError
            return int(z)
        if typ == "jahr":
            j = int(_zahl(v))
            if not 1900 <= j <= 2200:
                raise ValueError
            return j
        if typ == "datum":
            return _datum(v)
    except (ValueError, TypeError):
        fehler.append(f"{ort}: '{v}' ist kein gültiger Wert vom Typ {typ}")
        return None
    t = str(v).strip()
    if isinstance(v, float) and v == int(v):
        t = str(int(v))                                # 30.0 -> "30" bei Textfeldern
    if typ == "auswahl":
        erlaubt = vokabular.get(auswahl, VOKABULAR[auswahl])
        if t not in erlaubt:
            treffer = [e for e in erlaubt if e.lower() == t.lower()]
            if treffer:
                warnungen.append(f"{ort}: '{t}' -> '{treffer[0]}' (Schreibweise des Vokabulars)")
                return treffer[0]
            meldung = f"{ort}: '{t}' nicht im Vokabular {auswahl} ({' | '.join(erlaubt)})"
            if auswahl in OFFENE_LISTEN:
                warnungen.append(meldung + " - zusätzlicher Wert, vom Template erlaubt")
            else:
                fehler.append(meldung)
                return None
        return t
    if typ == "liste":
        return [x.strip() for x in t.split(";") if x.strip()]
    if typ == "url" and not re.match(r"https?://", t):
        fehler.append(f"{ort}: '{t}' ist keine URL (http/https)")
        return None
    if typ == "email" and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", t):
        fehler.append(f"{ort}: '{t}' ist keine E-Mail-Adresse")
        return None
    unbekannt = platzhalter_in(t) - set(PLATZHALTER)
    if unbekannt:
        fehler.append(f"{ort}: unbekannte Platzhalter {', '.join('{' + u + '}' for u in sorted(unbekannt))}")
    return t


def _kurz(liste, n=8):
    return ", ".join(liste[:n]) + (f" ... (+{len(liste) - n})" if len(liste) > n else "")


def pruefe(roh, vokabular=None, basis_ordner=None):
    """Rohdaten prüfen und in ein Modell ueberfuehren.

    vokabular: Listen aus dem AASX-Template ({name: [werte]}), ergänzt/ueberschreibt VOKABULAR.
    Rückgabe: (modell, fehler, warnungen). Bei Fehlern ist das Modell nicht verwendbar.
    """
    vok = {**VOKABULAR, **(vokabular or {})}
    fehler, warnungen = list(roh["fehler"]), []

    # -- Blatt EPD
    epd = {}
    for f in EPD_FELDER:
        if f.schluessel not in roh["epd"]:
            fehler.append(f"Blatt EPD: Feld '{f.schluessel}' ({f.label}) fehlt - Zeile gelöscht?")
            continue
        v, ort = roh["epd"][f.schluessel]
        ort = f"{ort} ({f.label})"
        if _leer(v):
            if f.pflicht:
                fehler.append(f"{ort}: Pflichtfeld ist leer")
            epd[f.schluessel] = None
            continue
        epd[f.schluessel] = _pruefe_wert(f.typ, f.auswahl, v, ort, fehler, warnungen, vok)
    for schl in roh["epd"]:
        if schl not in FELD:
            warnungen.append(f"Blatt EPD: unbekanntes Feld '{schl}' wird ignoriert")

    version = str(epd.get("vorlagen_version") or "")
    if version and version.split(".")[0] != VORLAGEN_VERSION.split(".")[0]:
        fehler.append(f"Vorlagenversion {version} passt nicht zu diesem Skript (v{VORLAGEN_VERSION})")
    if epd.get("ausgabedatum") and epd.get("gueltig_bis") and epd["gueltig_bis"] <= epd["ausgabedatum"]:
        fehler.append("Blatt EPD: 'Gültig bis' liegt nicht nach dem Ausgabedatum")
    for teil, name in (("fe", "Funktionelle Einheit"), ("de", "Deklarierte Einheit")):
        if (epd.get(f"{teil}_menge") is None) != (epd.get(f"{teil}_einheit") is None):
            fehler.append(f"Blatt EPD: {name} braucht Menge UND Einheit (oder keins von beiden)")
    for s in epd.get("sprachen") or []:
        if not re.fullmatch(r"[a-z]{2}", s):
            warnungen.append(f"Blatt EPD: Sprache '{s}' ist kein ISO-639-1-Code (zwei Kleinbuchstaben)")

    # -- Listenblaetter typgerecht einlesen
    listen = {}
    for blatt, spalten in LISTEN.items():
        zeilen = []
        for z in roh["listen"].get(blatt, []):
            neu = {"_zeile": z["_zeile"], "_ort": f"{blatt}!Zeile {z['_zeile']}"}
            for s in spalten:
                if s.schluessel not in z:
                    continue
                ort = z["_zellen"][s.schluessel]
                v = z[s.schluessel]
                if _leer(v):
                    if s.pflicht:
                        fehler.append(f"{ort} ({s.kopf}): Pflichtwert fehlt")
                    neu[s.schluessel] = None
                else:
                    neu[s.schluessel] = _pruefe_wert(s.typ, s.auswahl, v, f"{ort} ({s.kopf})",
                                                     fehler, warnungen, vok)
            zeilen.append(neu)
        listen[blatt] = zeilen

    # -- Produkte
    produkte = listen["Produkte"]
    if not produkte:
        fehler.append("Blatt Produkte: kein Produkt eingetragen")
    artikel = [p["artikel"] for p in produkte if p.get("artikel")]
    doppelt = sorted({a for a in artikel if artikel.count(a) > 1})
    if doppelt:
        fehler.append(f"Blatt Produkte: Artikelnummern mehrfach: {', '.join(doppelt)}")
    refs = [p for p in produkte if p.get("referenz") == "ja"]
    if len(refs) != 1 and produkte:
        fehler.append(f"Blatt Produkte: genau ein Referenzprodukt nötig, gefunden {len(refs)}")
    ref = refs[0]["artikel"] if len(refs) == 1 else None
    for p in produkte:
        if p.get("referenz") == "ja" and p.get("herkunft") == "Faktor":
            fehler.append(f"{p['_ort']}: Referenzprodukt muss Werteherkunft 'Deklariert' haben")
        if p.get("masse_kg") is not None and p["masse_kg"] <= 0:
            fehler.append(f"{p['_ort']}: Masse muss größer 0 sein")
        if p.get("masse_kg") is not None and p["masse_kg"] > 1000:
            warnungen.append(f"{p['_ort']}: Masse {p['masse_kg']} kg - Einheit kg statt g?")
        for s in ("bild",):
            if p.get(s) and basis_ordner and not (Path(basis_ordner) / p[s]).is_file():
                warnungen.append(f"{p['_ort']}: Datei '{p[s]}' nicht gefunden (relativ zur Excel)")
        for s in ("serie", "produktname", "beschreibung"):
            if p.get(s):
                unbekannt = platzhalter_in(p[s]) - set(PLATZHALTER)
                if unbekannt:
                    fehler.append(f"{p['_ort']}: unbekannte Platzhalter in {s}: {unbekannt}")
    if epd.get("scope_typ") == "Instance" and len(produkte) > 1:
        warnungen.append("scopeType Instance, aber mehrere Produkte - für Produktfamilien 'Class' verwenden")
    if epd.get("scope_typ") == "Class" and len(produkte) == 1:
        warnungen.append("scopeType Class, aber nur ein Produkt - für Einzelprodukt-EPDs 'Instance' verwenden")
    if epd.get("epd_dokument") and basis_ordner and not (Path(basis_ordner) / epd["epd_dokument"]).is_file():
        warnungen.append(f"Blatt EPD: EPD-PDF '{epd['epd_dokument']}' nicht gefunden (relativ zur Excel)")
    if not epd.get("epd_dokument"):
        warnungen.append("Blatt EPD: kein EPD-PDF angegeben - epdDocument (im Template Pflicht) entfällt")

    # -- Module
    module = [m for m in listen["Module"] if m.get("stageCode")]
    codes = [m["stageCode"] for m in module]
    for c in sorted({c for c in codes if codes.count(c) > 1}):
        fehler.append(f"Blatt Module: stageCode '{c}' mehrfach")
    if SUMMEN_STUFE in codes:
        fehler.append(f"Blatt Module: '{SUMMEN_STUFE}' ist kein Modul (Summenwerte nur im Blatt Ergebnisse)")
    status = {m["stageCode"]: m.get("stageStatus") for m in module}
    if not module:
        fehler.append("Blatt Module: keine Lebenszyklusmodule eingetragen")

    # -- Indikatoren
    indikatoren = {}
    for i in listen["Indikatoren"]:
        if not i.get("code"):
            continue
        if i["code"] in indikatoren:
            fehler.append(f"{i['_ort']}: indicatorCode '{i['code']}' mehrfach")
        indikatoren[i["code"]] = i
    reg = vok.get("indicatorCode")
    if reg:
        extra = sorted(set(indikatoren) - set(reg))
        if extra:
            warnungen.append("indicatorCodes außerhalb des Template-Vokabulars (vom Template zugelassen, "
                             f"'where required by the applicable PCR'): {_kurz(extra)}")

    # -- Ergebnisse
    ergebnisse = {}
    for e in listen["Ergebnisse"]:
        a, c, s = e.get("artikel"), e.get("code"), e.get("stage")
        if not (a and c and s) or e.get("wert") is None:
            continue
        if a not in artikel:
            fehler.append(f"{e['_ort']}: Artikel '{a}' steht nicht im Blatt Produkte")
        if c not in indikatoren:
            fehler.append(f"{e['_ort']}: indicatorCode '{c}' steht nicht im Blatt Indikatoren")
        if s != SUMMEN_STUFE and s not in status:
            fehler.append(f"{e['_ort']}: stageCode '{s}' steht nicht im Blatt Module")
        elif status.get(s) == "NotAssessed":
            fehler.append(f"{e['_ort']}: Wert für Modul '{s}', das als NotAssessed markiert ist")
        if (a, c, s) in ergebnisse:
            fehler.append(f"{e['_ort']}: Wert für {a} / {c} / {s} mehrfach")
        ergebnisse[(a, c, s)] = e
    for s, st in status.items():
        if st == "Assessed" and ref and not any(k[0] == ref and k[2] == s for k in ergebnisse):
            warnungen.append(f"Modul {s} ist Assessed, hat aber keinen einzigen Wert für {ref}")

    # -- Faktoren
    faktoren = {}
    for f in listen["Faktoren"]:
        a, c, s = f.get("artikel"), f.get("code"), f.get("stage")
        if not (a and c and s) or f.get("faktor") is None:
            continue
        if a not in artikel:
            warnungen.append(f"{f['_ort']}: Faktor für '{a}' (nicht im Blatt Produkte) wird ignoriert")
            continue
        if c not in indikatoren:
            fehler.append(f"{f['_ort']}: indicatorCode '{c}' steht nicht im Blatt Indikatoren")
        if s != SUMMEN_STUFE and s not in status:
            fehler.append(f"{f['_ort']}: stageCode '{s}' steht nicht im Blatt Module")
        if (a, c, s) in faktoren:
            fehler.append(f"{f['_ort']}: Faktor für {a} / {c} / {s} mehrfach")
        faktoren[(a, c, s)] = f["faktor"]

    for p in produkte:
        a = p.get("artikel")
        eigene = [k for k in ergebnisse if k[0] == a]
        if p.get("herkunft") == "Deklariert" and not eigene:
            fehler.append(f"{p['_ort']}: Werteherkunft Deklariert, aber keine Werte im Blatt Ergebnisse")
        if p.get("herkunft") == "Faktor":
            if eigene:
                fehler.append(f"{p['_ort']}: Werteherkunft Faktor, aber eigene Werte im Blatt Ergebnisse")
            if ref:
                fehlt = [f"{c}/{s}" for (r, c, s) in ergebnisse if r == ref and (a, c, s) not in faktoren]
                if fehlt:
                    fehler.append(f"{p['_ort']}: {len(fehlt)} Faktoren fehlen für {a}: {_kurz(fehlt)}")

    # -- Standorte, Normen
    for s in listen["Standorte"]:
        if s.get("land") and not re.fullmatch(r"[A-Z]{2}", s["land"]):
            warnungen.append(f"{s['_ort']}: countryCode '{s['land']}' ist kein ISO-3166-1-alpha-2-Code")
    normen = [n["bezeichnung"] for n in listen["Normen_PCR"] if n.get("art") == "Norm" and n.get("bezeichnung")]
    regeln = [n["bezeichnung"] for n in listen["Normen_PCR"] if n.get("art") in ("PCR", "PSR")
              and n.get("bezeichnung")]
    if not normen:
        fehler.append("Blatt Normen_PCR: mindestens eine Norm nötig (standardsCompliance ist Pflicht)")
    if not regeln:
        fehler.append("Blatt Normen_PCR: mindestens eine PCR nötig (productCategoryRules ist Pflicht)")

    modell = {
        "epd": epd, "produkte": produkte, "referenz": ref, "module": module, "indikatoren": indikatoren,
        "ergebnisse": {k: (v["wert"], v.get("einheit_quelle")) for k, v in ergebnisse.items()},
        "faktoren": faktoren, "standorte": listen["Standorte"], "normen": normen, "regeln": regeln,
    }
    return modell, fehler, list(dict.fromkeys(warnungen))


def ausgeben(fehler, warnungen, datei=None):
    datei = datei or sys.stdout
    if warnungen:
        print(f"\nWarnungen ({len(warnungen)}):", file=datei)
        for w in warnungen:
            print(f"  - {w}", file=datei)
    if fehler:
        print(f"\nFEHLER ({len(fehler)}):", file=datei)
        for f in fehler:
            print(f"  x {f}", file=datei)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--leer", metavar="XLSX", help="leere Vorlage schreiben")
    g.add_argument("--pruefe", metavar="XLSX", help="ausgefüllte Excel prüfen")
    args = ap.parse_args()
    if args.leer:
        schreibe_arbeitsmappe(args.leer)
        print(f"Vorlage v{VORLAGEN_VERSION} geschrieben: {args.leer}")
        return
    roh = lese_arbeitsmappe(args.pruefe)
    modell, fehler, warnungen = pruefe(roh, basis_ordner=Path(args.pruefe).resolve().parent)
    ausgeben(fehler, warnungen)
    n = len(modell["produkte"])
    print(f"\n{args.pruefe}: {n} Produkte, {len(modell['ergebnisse'])} Werte, {len(modell['faktoren'])} Faktoren - "
          + ("NICHT in Ordnung" if fehler else "in Ordnung"))
    sys.exit(2 if fehler else 0)


if __name__ == "__main__":
    main()
